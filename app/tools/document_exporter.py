import os
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from app.core.config import BASE_DIR
from app.agents.tailoring.schema import TailoredResume, TailoredCoverLetter

EXPORTS_DIR = BASE_DIR / "data" / "exports"
PROFILE_PATH = BASE_DIR / "data" / "profile" / "profile.json"


def ensure_export_dir() -> Path:
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    return EXPORTS_DIR


def get_profile_info() -> dict:
    """Loads candidate contact info and education from profile.json."""
    if PROFILE_PATH.exists():
        with open(PROFILE_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def set_cell_margins(cell, top=0, bottom=0, left=0, right=0):
    """Sets padding for borderless table cells."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_section_header(doc: Document, title: str):
    """Adds a bold section heading with a clean horizontal bottom border."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    
    run = p.add_run(title.upper())
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(30, 41, 59)  # Slate dark

    # Add bottom border XML to paragraph
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                     r'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="94A3B8"/>'
                     r'</w:pBdr>')
    pPr.append(pBdr)


def add_entry_header_table(doc: Document, left_title: str, left_subtitle: str, right_top: str, right_bottom: str):
    """
    Creates a borderless 2-column table:
    Left: Bold Title, Subtitle
    Right: Right-aligned Date, Location
    """
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # 7.5 inches printable width (8.5 - 2*0.5 margins)
    table.columns[0].width = Inches(5.3)
    table.columns[1].width = Inches(2.2)

    # Row 1: Title (Left) and Date/Period (Right)
    r1c1 = table.cell(0, 0).paragraphs[0]
    r1c1.paragraph_format.space_after = Pt(0)
    run1 = r1c1.add_run(left_title)
    run1.bold = True
    run1.font.name = "Arial"
    run1.font.size = Pt(10)

    r1c2 = table.cell(0, 1).paragraphs[0]
    r1c2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1c2.paragraph_format.space_after = Pt(0)
    run2 = r1c2.add_run(right_top)
    run2.bold = True
    run2.font.name = "Arial"
    run2.font.size = Pt(9.5)

    # Row 2: Subtitle (Left) and Location (Right)
    r2c1 = table.cell(1, 0).paragraphs[0]
    r2c1.paragraph_format.space_after = Pt(2)
    run3 = r2c1.add_run(left_subtitle)
    run3.italic = True
    run3.font.name = "Arial"
    run3.font.size = Pt(9.5)

    r2c2 = table.cell(1, 1).paragraphs[0]
    r2c2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2c2.paragraph_format.space_after = Pt(2)
    run4 = r2c2.add_run(right_bottom)
    run4.font.name = "Arial"
    run4.font.size = Pt(9.5)

    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=0, bottom=0, left=0, right=0)


def export_resume_docx(resume: TailoredResume) -> Path:
    """Exports a professional, ATS-optimized 1-page format resume in Word (.docx)."""
    ensure_export_dir()
    profile = get_profile_info()
    file_path = EXPORTS_DIR / f"{resume.target_job_id}_resume.docx"

    doc = Document()

    # Set 0.5 inch margins all around
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.5)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.5)
        s.right_margin = Inches(0.5)

    # --- 1. HEADER (Candidate Contact Info) ---
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_after = Pt(2)
    name_run = name_p.add_run(profile.get("full_name") or "Candidate")
    name_run.bold = True
    name_run.font.name = "Arial"
    name_run.font.size = Pt(20)
    name_run.font.color.rgb = RGBColor(15, 23, 42)

    # Contact line - safely check each field to prevent NoneType errors
    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_p.paragraph_format.space_after = Pt(8)
    
    contact_parts = []
    if profile.get("phone"):
        contact_parts.append(str(profile["phone"]).strip())
    if profile.get("email"):
        contact_parts.append(str(profile["email"]).strip())
    if profile.get("location"):
        contact_parts.append(str(profile["location"]).strip())
    if profile.get("linkedin_url"):
        clean_li = str(profile["linkedin_url"]).replace("https://", "").replace("http://", "").rstrip("/")
        contact_parts.append(clean_li)
    if profile.get("github_url"):
        clean_gh = str(profile["github_url"]).replace("https://", "").replace("http://", "").rstrip("/")
        contact_parts.append(clean_gh)

    if not contact_parts:
        contact_parts = [profile.get("email") or "applicant@example.com"]

    contact_run = contact_p.add_run("  |  ".join(contact_parts))
    contact_run.font.name = "Arial"
    contact_run.font.size = Pt(9)
    contact_run.font.color.rgb = RGBColor(71, 85, 105)

    # --- 2. PROFESSIONAL SUMMARY ---
    if resume.professional_summary:
        add_section_header(doc, "Professional Summary")
        sum_p = doc.add_paragraph(resume.professional_summary)
        sum_p.paragraph_format.space_after = Pt(4)
        sum_p.paragraph_format.line_spacing = 1.15
        for run in sum_p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(9.5)

    # --- 3. TECHNICAL SKILLS ---
    if resume.highlighted_skills:
        add_section_header(doc, "Technical Skills")
        skills_p = doc.add_paragraph()
        skills_p.paragraph_format.space_after = Pt(4)
        s_bold = skills_p.add_run("Core & Matching Competencies: ")
        s_bold.bold = True
        s_bold.font.name = "Arial"
        s_bold.font.size = Pt(9.5)
        s_text = skills_p.add_run(", ".join(resume.highlighted_skills))
        s_text.font.name = "Arial"
        s_text.font.size = Pt(9.5)

    # --- 4. PROFESSIONAL EXPERIENCE ---
    if resume.experience:
        add_section_header(doc, "Professional Experience")
        for exp in resume.experience:
            add_entry_header_table(
                doc,
                left_title=exp.role,
                left_subtitle=exp.company,
                right_top=exp.period,
                right_bottom="Remote"
            )
            for b in exp.bullets:
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_before = Pt(0)
                bp.paragraph_format.space_after = Pt(2)
                bp.paragraph_format.line_spacing = 1.1
                brun = bp.add_run(b.text)
                brun.font.name = "Arial"
                brun.font.size = Pt(9.5)

    # --- 5. KEY TECHNICAL PROJECTS ---
    if resume.projects:
        add_section_header(doc, "Key Technical Projects")
        for proj in resume.projects:
            tech_str = f"Technologies: {', '.join(proj.technologies)}" if proj.technologies else ""
            add_entry_header_table(
                doc,
                left_title=proj.title,
                left_subtitle=tech_str,
                right_top="Personal Project",
                right_bottom="GitHub"
            )
            for b in proj.bullets:
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_before = Pt(0)
                bp.paragraph_format.space_after = Pt(2)
                bp.paragraph_format.line_spacing = 1.1
                brun = bp.add_run(b.text)
                brun.font.name = "Arial"
                brun.font.size = Pt(9.5)

    # --- 6. EDUCATION ---
    education_list = profile.get("education", [])
    if education_list:
        add_section_header(doc, "Education")
        for edu in education_list:
            if not isinstance(edu, dict):
                continue
            deg = edu.get("degree") or "Degree"
            field = edu.get("field_of_study") or ""
            degree_str = f"{deg} in {field}".strip() if field else deg
            start = str(edu.get("start_year") or "")
            end = str(edu.get("end_year") or "Present")
            dates_str = f"{start} - {end}".strip(" -")
            gpa_str = f"GPA: {edu.get('gpa')}" if edu.get("gpa") else ""
            add_entry_header_table(
                doc,
                left_title=edu.get("institution") or "University",
                left_subtitle=f"{degree_str} ({gpa_str})" if gpa_str else degree_str,
                right_top=dates_str,
                right_bottom=""
            )

    doc.save(str(file_path))
    return file_path


def export_cover_letter_docx(letter: TailoredCoverLetter) -> Path:
    """Exports a clean, executive cover letter formatted for Word."""
    ensure_export_dir()
    profile = get_profile_info()
    file_path = EXPORTS_DIR / f"{letter.target_job_id}_cover_letter.docx"

    doc = Document()

    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    # Header
    name_p = doc.add_paragraph()
    name_p.paragraph_format.space_after = Pt(2)
    name_run = name_p.add_run(profile.get("full_name") or "Candidate")
    name_run.bold = True
    name_run.font.name = "Arial"
    name_run.font.size = Pt(16)

    contact_p = doc.add_paragraph()
    contact_p.paragraph_format.space_after = Pt(16)
    
    c_parts = []
    if profile.get("email"):
        c_parts.append(str(profile["email"]).strip())
    if profile.get("phone"):
        c_parts.append(str(profile["phone"]).strip())
    if profile.get("location"):
        c_parts.append(str(profile["location"]).strip())
    
    contact_str = "  |  ".join(c_parts) if c_parts else (profile.get("email") or "applicant@example.com")
    contact_run = contact_p.add_run(contact_str)
    contact_run.font.name = "Arial"
    contact_run.font.size = Pt(9.5)
    contact_run.font.color.rgb = RGBColor(100, 116, 139)

    # Recipient
    recip_p = doc.add_paragraph()
    recip_p.paragraph_format.space_after = Pt(12)
    r_run = recip_p.add_run(f"Hiring Team\n{letter.company_name}\nRe: Application for {letter.role_title}")
    r_run.bold = True
    r_run.font.name = "Arial"
    r_run.font.size = Pt(10)

    # Salutation
    sal_p = doc.add_paragraph(letter.salutation)
    sal_p.paragraph_format.space_after = Pt(8)
    if sal_p.runs:
        sal_p.runs[0].font.name = "Arial"
        sal_p.runs[0].font.size = Pt(10)

    # Opening
    op_p = doc.add_paragraph(letter.opening_paragraph)
    op_p.paragraph_format.space_after = Pt(8)
    op_p.paragraph_format.line_spacing = 1.15
    for r in op_p.runs:
        r.font.name = "Arial"
        r.font.size = Pt(10)

    # Body
    for para in letter.body_paragraphs:
        bp = doc.add_paragraph(para)
        bp.paragraph_format.space_after = Pt(8)
        bp.paragraph_format.line_spacing = 1.15
        for r in bp.runs:
            r.font.name = "Arial"
            r.font.size = Pt(10)

    # Closing
    cp = doc.add_paragraph(letter.closing_paragraph)
    cp.paragraph_format.space_after = Pt(14)
    cp.paragraph_format.line_spacing = 1.15
    for r in cp.runs:
        r.font.name = "Arial"
        r.font.size = Pt(10)

    # Sign-off
    so_p = doc.add_paragraph(letter.sign_off)
    for r in so_p.runs:
        r.font.name = "Arial"
        r.font.size = Pt(10)

    doc.save(str(file_path))
    return file_path


def export_resume_markdown(resume: TailoredResume) -> Path:
    ensure_export_dir()
    file_path = EXPORTS_DIR / f"{resume.target_job_id}_resume.md"
    skills_str = ", ".join(resume.highlighted_skills)
    
    exp_sections = []
    for exp in resume.experience:
        bullets = "\n".join([f"- {b.text}" for b in exp.bullets])
        exp_sections.append(f"### {exp.role} — {exp.company}\n*{exp.period}*\n{bullets}")

    proj_sections = []
    for proj in resume.projects:
        tech = ", ".join(proj.technologies)
        bullets = "\n".join([f"- {b.text}" for b in proj.bullets])
        proj_sections.append(f"### {proj.title} ({tech})\n{bullets}")

    sections = [
        f"# {resume.target_role} — Tailored for {resume.target_company}",
        f"## Professional Summary\n{resume.professional_summary}",
        f"## Highlighted Technical Skills\n{skills_str}"
    ]
    if exp_sections:
        sections.append("## Professional Experience\n" + "\n\n".join(exp_sections))
    if proj_sections:
        sections.append("## Key Technical Projects\n" + "\n\n".join(proj_sections))

    content = "\n\n".join(sections)

    with open(file_path, "w", encoding="utf-8") as f:
        f.write(content.strip())
    return file_path


def export_cover_letter_markdown(letter: TailoredCoverLetter) -> Path:
    ensure_export_dir()
    file_path = EXPORTS_DIR / f"{letter.target_job_id}_cover_letter.md"
    body = "\n\n".join(letter.body_paragraphs)
    content = f"""# Cover Letter: {letter.role_title} at {letter.company_name}

{letter.salutation}

{letter.opening_paragraph}

{body}

{letter.closing_paragraph}

{letter.sign_off}
"""
    with open(file_path, "w", encoding="utf-8") as f:
        f.write(content.strip())
    return file_path


def build_resume_html(resume: TailoredResume) -> str:
    """Generates a clean, ATS-optimized, single-page HTML layout for PDF rendering."""
    import html
    profile = get_profile_info()
    full_name = html.escape(profile.get("full_name") or "Candidate")

    contact_parts = []
    if profile.get("phone"):
        contact_parts.append(html.escape(str(profile["phone"]).strip()))
    if profile.get("email"):
        contact_parts.append(html.escape(str(profile["email"]).strip()))
    if profile.get("location"):
        contact_parts.append(html.escape(str(profile["location"]).strip()))
    if profile.get("linkedin_url"):
        clean_li = html.escape(str(profile["linkedin_url"]).replace("https://", "").replace("http://", "").rstrip("/"))
        contact_parts.append(clean_li)
    if profile.get("github_url"):
        clean_gh = html.escape(str(profile["github_url"]).replace("https://", "").replace("http://", "").rstrip("/"))
        contact_parts.append(clean_gh)

    contact_line = " &bull; ".join(contact_parts) if contact_parts else (profile.get("email") or "applicant@example.com")

    summary_html = ""
    if resume.professional_summary:
        summary_html = f"""
        <div class="section-title">Professional Summary</div>
        <p class="summary">{html.escape(resume.professional_summary)}</p>
        """

    skills_html = ""
    if resume.highlighted_skills:
        skills_str = html.escape(", ".join(resume.highlighted_skills))
        skills_html = f"""
        <div class="section-title">Technical Skills</div>
        <p class="skills-line"><span class="skills-bold">Core & Matching Competencies:</span> {skills_str}</p>
        """

    exp_html = ""
    if resume.experience:
        exp_items = []
        for exp in resume.experience:
            role = html.escape(exp.role)
            company = html.escape(exp.company)
            period = html.escape(exp.period or "Present")
            bullets = "".join([f"<li>{html.escape(b.text)}</li>" for b in exp.bullets])
            exp_items.append(f"""
            <div class="item">
              <div class="item-header">
                <span class="item-title">{role}</span>
                <span class="item-date">{period}</span>
              </div>
              <div class="item-sub">{company}</div>
              <ul class="bullets">{bullets}</ul>
            </div>
            """)
        exp_html = f"""
        <div class="section-title">Professional Experience</div>
        {"".join(exp_items)}
        """

    proj_html = ""
    if resume.projects:
        proj_items = []
        for proj in resume.projects:
            title = html.escape(proj.title)
            tech = html.escape(", ".join(proj.technologies)) if proj.technologies else ""
            bullets = "".join([f"<li>{html.escape(b.text)}</li>" for b in proj.bullets])
            proj_items.append(f"""
            <div class="item">
              <div class="item-header">
                <span class="item-title">{title}</span>
                <span class="item-date">Personal Project</span>
              </div>
              {f'<div class="item-sub">Technologies: {tech}</div>' if tech else ''}
              <ul class="bullets">{bullets}</ul>
            </div>
            """)
        proj_html = f"""
        <div class="section-title">Key Technical Projects</div>
        {"".join(proj_items)}
        """

    edu_html = ""
    education_list = profile.get("education", [])
    if education_list:
        edu_items = []
        for edu in education_list:
            if not isinstance(edu, dict):
                continue
            inst = html.escape(edu.get("institution") or "University")
            deg = html.escape(edu.get("degree") or "Degree")
            field = html.escape(edu.get("field_of_study") or "")
            degree_str = f"{deg} in {field}".strip() if field else deg
            start = str(edu.get("start_year") or "")
            end = str(edu.get("end_year") or "Present")
            dates_str = html.escape(f"{start} - {end}".strip(" -"))
            gpa = edu.get("gpa")
            gpa_str = f" &bull; GPA: {html.escape(str(gpa))}" if gpa else ""
            edu_items.append(f"""
            <div class="item">
              <div class="item-header">
                <span class="item-title">{inst}</span>
                <span class="item-date">{dates_str}</span>
              </div>
              <div class="item-sub">{degree_str}{gpa_str}</div>
            </div>
            """)
        edu_html = f"""
        <div class="section-title">Education</div>
        {"".join(edu_items)}
        """

    return f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  @page {{
    size: A4;
    margin: 12mm 15mm;
  }}
  * {{
    box-sizing: border-box;
    margin: 0;
    padding: 0;
  }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, 'Helvetica Neue', Arial, sans-serif;
    color: #0f172a;
    background: #ffffff;
    line-height: 1.35;
    font-size: 10pt;
  }}
  .header {{
    text-align: center;
    margin-bottom: 10px;
  }}
  .name {{
    font-size: 18pt;
    font-weight: 700;
    color: #0f172a;
    letter-spacing: -0.3px;
    margin-bottom: 3px;
  }}
  .contact {{
    font-size: 8.5pt;
    color: #475569;
  }}
  .section-title {{
    font-size: 9.5pt;
    font-weight: 700;
    text-transform: uppercase;
    letter-spacing: 0.6px;
    color: #0f172a;
    border-bottom: 1px solid #cbd5e1;
    padding-bottom: 2px;
    margin-top: 10px;
    margin-bottom: 5px;
  }}
  .summary {{
    font-size: 9pt;
    color: #334155;
    margin-bottom: 4px;
    line-height: 1.35;
  }}
  .skills-line {{
    font-size: 9pt;
    color: #334155;
    margin-bottom: 4px;
  }}
  .skills-bold {{
    font-weight: 600;
    color: #0f172a;
  }}
  .item {{
    margin-bottom: 6px;
  }}
  .item-header {{
    display: flex;
    justify-content: space-between;
    align-items: baseline;
    font-size: 9.5pt;
  }}
  .item-title {{
    font-weight: 700;
    color: #0f172a;
  }}
  .item-date {{
    font-size: 8.5pt;
    font-weight: 600;
    color: #64748b;
  }}
  .item-sub {{
    font-size: 8.5pt;
    color: #475569;
    font-style: italic;
    margin-bottom: 2px;
  }}
  ul.bullets {{
    padding-left: 16px;
    margin-top: 2px;
  }}
  ul.bullets li {{
    font-size: 9pt;
    color: #334155;
    margin-bottom: 2px;
    line-height: 1.3;
  }}
</style>
</head>
<body>
  <div class="header">
    <div class="name">{full_name}</div>
    <div class="contact">{contact_line}</div>
  </div>
  {summary_html}
  {skills_html}
  {exp_html}
  {proj_html}
  {edu_html}
</body>
</html>
"""


def build_cover_letter_html(letter: TailoredCoverLetter) -> str:
    """Generates an executive, formal business letter layout for PDF rendering."""
    import html
    profile = get_profile_info()
    full_name = html.escape(profile.get("full_name") or "Candidate")
    email = html.escape(str(profile.get("email") or ""))
    phone = html.escape(str(profile.get("phone") or ""))
    location = html.escape(str(profile.get("location") or ""))

    contact_parts = [p for p in [email, phone, location] if p]
    contact_line = " &bull; ".join(contact_parts)

    paragraphs_html = "".join([f"<p>{html.escape(p)}</p>" for p in letter.body_paragraphs])

    return f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  @page {{
    size: A4;
    margin: 20mm 22mm;
  }}
  * {{
    box-sizing: border-box;
    margin: 0;
    padding: 0;
  }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, 'Helvetica Neue', Arial, sans-serif;
    color: #0f172a;
    background: #ffffff;
    line-height: 1.5;
    font-size: 10.5pt;
  }}
  .header {{
    border-bottom: 1px solid #e2e8f0;
    padding-bottom: 12px;
    margin-bottom: 20px;
  }}
  .name {{
    font-size: 18pt;
    font-weight: 700;
    color: #0f172a;
    margin-bottom: 2px;
  }}
  .contact {{
    font-size: 9pt;
    color: #64748b;
  }}
  .recipient {{
    margin-bottom: 20px;
    font-size: 10pt;
    color: #1e293b;
    line-height: 1.4;
  }}
  .recipient-title {{
    font-weight: 700;
    color: #0f172a;
  }}
  .content p {{
    margin-bottom: 12px;
    color: #334155;
    text-align: justify;
  }}
  .signoff {{
    margin-top: 20px;
    color: #0f172a;
    line-height: 1.4;
  }}
  .signoff-name {{
    font-weight: 700;
  }}
</style>
</head>
<body>
  <div class="header">
    <div class="name">{full_name}</div>
    <div class="contact">{contact_line}</div>
  </div>

  <div class="recipient">
    <div class="recipient-title">Hiring Team</div>
    <div>{html.escape(letter.company_name)}</div>
    <div><strong>Re:</strong> Application for {html.escape(letter.role_title)}</div>
  </div>

  <div class="content">
    <p>{html.escape(letter.salutation)}</p>
    <p>{html.escape(letter.opening_paragraph)}</p>
    {paragraphs_html}
    <p>{html.escape(letter.closing_paragraph)}</p>
  </div>

  <div class="signoff">
    <div>Sincerely,</div>
    <div class="signoff-name">{full_name}</div>
  </div>
</body>
</html>
"""


async def export_resume_pdf(resume: TailoredResume) -> Path:
    """Renders pixel-perfect ATS resume PDF using headless Chromium."""
    ensure_export_dir()
    file_path = EXPORTS_DIR / f"{resume.target_job_id}_resume.pdf"
    html_content = build_resume_html(resume)

    from playwright.async_api import async_playwright
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        page = await browser.new_page()
        await page.set_content(html_content, wait_until="networkidle")
        await page.pdf(
            path=str(file_path),
            format="A4",
            print_background=True,
            margin={"top": "10mm", "bottom": "10mm", "left": "12mm", "right": "12mm"}
        )
        await browser.close()
    return file_path


async def export_cover_letter_pdf(letter: TailoredCoverLetter) -> Path:
    """Renders executive business cover letter PDF using headless Chromium."""
    ensure_export_dir()
    file_path = EXPORTS_DIR / f"{letter.target_job_id}_cover_letter.pdf"
    html_content = build_cover_letter_html(letter)

    from playwright.async_api import async_playwright
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        page = await browser.new_page()
        await page.set_content(html_content, wait_until="networkidle")
        await page.pdf(
            path=str(file_path),
            format="A4",
            print_background=True,
            margin={"top": "15mm", "bottom": "15mm", "left": "18mm", "right": "18mm"}
        )
        await browser.close()
    return file_path
